#!/usr/bin/env python3
"""OTscope detection regression suite.

Builds synthetic captures (pure stdlib, see pcapgen.py / scenarios.py), runs
`otscope.py --scan --offline` on each in a throwaway workspace, and checks the
JSON / CSV / DOCX / SVG output against per-scenario expectations:

  * benign scenarios must not raise intrusion-class or known false-positive findings
  * attack scenarios must still raise the findings their behavior warrants

Requirements: Python 3.8+, tshark on PATH, and python-docx (already an OTscope
runtime dependency; DOCX checks are skipped if it is missing).  Nothing is
written inside the repository — workspaces go to a temp directory and the
user's ~/.ot_pcap_analyzer.conf is never touched.

Usage:
    python tests/run_regression.py                 # run everything
    python tests/run_regression.py --only av_camera_voip benign_scada_master
    python tests/run_regression.py --keep          # keep the temp workspaces for inspection
"""

from __future__ import annotations

import argparse
import csv
import json
import os
import shutil
import subprocess
import sys
import tempfile
import xml.dom.minidom
from pathlib import Path
from typing import Any, Callable, Dict, List, Optional

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))

from scenarios import SCENARIOS  # noqa: E402

REPO = HERE.parent
DEFAULT_OTSCOPE = REPO / "otscope" / "src" / "otscope.py"
DNP3_SAMPLE = next((path for path in (REPO / "otscope" / "src" / "pcaps" / "dnp3_challenge.pcap",
                                      REPO / "otscope" / "pcaps" / "dnp3_challenge.pcap") if path.exists()),
                   REPO / "otscope" / "src" / "pcaps" / "dnp3_challenge.pcap")
BENIGN_FORBIDDEN_SEVERITIES = {"INTRUSION", "CORRELATION"}
BAND_ORDER = ["MINIMAL", "LOW", "MODERATE", "HIGH", "SEVERE"]


class Result:
    def __init__(self, name: str) -> None:
        self.name = name
        self.failures: List[str] = []
        self.summary = ""

    def check(self, ok: bool, message: str) -> None:
        if not ok:
            self.failures.append(message)


class RunOutput:
    def __init__(self, workspace: Path, proc: subprocess.CompletedProcess) -> None:
        self.workspace = workspace
        self.returncode = proc.returncode
        self.log = (proc.stdout or "") + (proc.stderr or "")
        out = workspace / "output"
        jsons = sorted(out.glob("OT_PCAP_Analysis_*.json"))
        self.report: Dict[str, Any] = json.loads(jsons[-1].read_text(encoding="utf-8")) if jsons else {}
        allow = sorted(out.glob("OT_Flow_Allowlist_*.csv"))
        self.allowlist: List[Dict[str, str]] = []
        if allow:
            with allow[-1].open(encoding="utf-8", newline="") as handle:
                self.allowlist = list(csv.DictReader(handle))
        self.svgs = sorted(out.glob("*.svg"))
        docs = sorted(out.glob("*.docx"))
        self.docx_headings: Optional[List[str]] = None
        self.docx_text = ""
        self.unreadable_header_cells = 0
        if docs:
            try:
                import docx  # type: ignore
                from docx.oxml.ns import qn  # type: ignore
                document = docx.Document(str(docs[-1]))
                self.docx_headings = [p.text for p in document.paragraphs if p.style.name.startswith("Heading")]
                parts = [p.text for p in document.paragraphs]
                for table in document.tables:
                    for row in table.rows:
                        parts.extend(cell.text for cell in row.cells)
                    # Header cells get the dark blue fill; their text must be white.
                    for cell in table.rows[0].cells if len(table.rows) else []:
                        shading = cell._tc.find(f"{qn('w:tcPr')}/{qn('w:shd')}")
                        if shading is None or shading.get(qn("w:fill")) != "2E5F8A":
                            continue
                        runs = [run for para in cell.paragraphs for run in para.runs]
                        if any(str(run.font.color.rgb) != "FFFFFF" for run in runs):
                            self.unreadable_header_cells += 1
                self.docx_text = "\n".join(parts)
            except ImportError:
                self.docx_headings = None

    # -- helpers --------------------------------------------------------------
    @property
    def findings(self) -> List[Dict[str, Any]]:
        return self.report.get("findings", [])

    def matching(self, contains: str, src: Optional[str] = None, dst: Optional[str] = None,
                 severity: Optional[str] = None) -> List[Dict[str, Any]]:
        hits = []
        for finding in self.findings:
            if contains.lower() not in finding["title"].lower():
                continue
            if src and src not in finding["source_ips"]:
                continue
            if dst and dst not in finding["destination_ips"]:
                continue
            if severity and finding["severity"] != severity:
                continue
            hits.append(finding)
        return hits

    def involving(self, contains: str, ip: str) -> List[Dict[str, Any]]:
        return [f for f in self.matching(contains) if ip in f["source_ips"] or ip in f["destination_ips"]]

    def band(self) -> str:
        return self.report.get("risk", {}).get("band", "?")

    def device(self, ip: str) -> Dict[str, Any]:
        return self.report.get("devices", {}).get(ip, {})


def run_otscope(otscope: Path, workspace: Path, captures: List[Path], loose: bool, site: Optional[str],
                home: Path) -> RunOutput:
    (workspace / "pcaps").mkdir(parents=True, exist_ok=True)
    shutil.copy2(otscope, workspace / "otscope.py")
    for capture in captures:
        target = workspace / capture.name if loose else workspace / "pcaps" / capture.name
        shutil.copy2(capture, target)
    cmd = [sys.executable, "otscope.py", "--scan", "pcaps", "--offline"]
    if site is not None:
        cmd += ["--site", site]
    env = dict(os.environ)
    env.update({"HOME": str(home), "USERPROFILE": str(home), "PYTHONIOENCODING": "utf-8"})
    proc = subprocess.run(cmd, cwd=str(workspace), env=env, capture_output=True, text=True,
                          encoding="utf-8", errors="replace", timeout=900)
    return RunOutput(workspace, proc)


# ---------------------------------------------------------------- expectations
def expect_common(r: Result, out: RunOutput) -> None:
    r.check(out.returncode == 0, f"otscope exited {out.returncode}")
    r.check(bool(out.report), "no JSON report produced")
    for svg in out.svgs:
        try:
            xml.dom.minidom.parse(str(svg))
        except Exception as exc:  # noqa: BLE001
            r.failures.append(f"{svg.name} is not valid XML ({exc})")
    r.check("Traceback" not in out.log, "Python traceback in run log")
    r.check("Word report generation failed" not in out.log, "Word report fell back to plain text")
    r.check(out.unreadable_header_cells == 0,
            f"{out.unreadable_header_cells} table header cell(s) have dark text on the dark header fill")


def expect_benign(r: Result, out: RunOutput, max_band: str, forbidden_titles: List[str]) -> None:
    for finding in out.findings:
        r.check(finding["severity"] not in BENIGN_FORBIDDEN_SEVERITIES,
                f"benign capture raised {finding['severity']}: {finding['title']}")
    for title in forbidden_titles:
        hits = out.matching(title)
        r.check(not hits, f"false positive: '{hits[0]['title']}'" if hits else "")
    band = out.band()
    r.check(band in BAND_ORDER and BAND_ORDER.index(band) <= BAND_ORDER.index(max_band),
            f"risk band {band} exceeds {max_band} (score {out.report.get('risk', {}).get('score')})")


def require(r: Result, out: RunOutput, contains: str, src: Optional[str] = None, dst: Optional[str] = None,
            severity: Optional[str] = None) -> None:
    hint = contains + (f" src={src}" if src else "") + (f" dst={dst}" if dst else "")
    r.check(bool(out.matching(contains, src, dst, severity)), f"missing expected finding: {hint}")


def forbid(r: Result, out: RunOutput, contains: str, src: Optional[str] = None, dst: Optional[str] = None) -> None:
    hint = contains + (f" src={src}" if src else "") + (f" dst={dst}" if dst else "")
    r.check(not out.matching(contains, src, dst), f"unexpected finding: {hint}")


def require_heading(r: Result, out: RunOutput, text: str) -> None:
    if out.docx_headings is None:
        return
    r.check(any(text.lower() in h.lower() for h in out.docx_headings), f"Word report has no heading containing '{text}'")


def require_text(r: Result, out: RunOutput, text: str) -> None:
    if out.docx_headings is None:
        return
    r.check(text in out.docx_text, f"Word report does not contain '{text}'")


NOISE = ["Unknown OT-adjacent traffic", "Legacy service observed"]


def check_benign_modbus_hmi(r: Result, out: RunOutput) -> None:
    expect_benign(r, out, "LOW", NOISE + ["register scanning", "replay", "Beaconing", "Dual-zone",
                                          "polling burst", "single-register polling", "Unexpected Modbus write source",
                                          "NTP observed"])
    require(r, out, "Modbus write operation observed", src="10.10.1.10", dst="10.10.1.20")
    forbid(r, out, "Modbus write operation observed", src="10.10.1.20")
    r.check(len(out.allowlist) == 3, f"flow allowlist has {len(out.allowlist)} rows, expected 3 real flows")
    r.check(out.report.get("site") == "benign_modbus_hmi",
            f"default site name is {out.report.get('site')!r}, expected the workspace folder name")
    r.check("Moved capture into workspace" in out.log, "loose capture next to otscope.py was not adopted into pcaps/")
    r.check("HMI" in out.device("10.10.1.10").get("inferred_role", ""), "HMI 10.10.1.10 not classified as HMI")
    r.check("Controller" in out.device("10.10.1.20").get("inferred_role", ""), "PLC 10.10.1.20 not classified as controller")
    r.check(49733 not in out.device("10.10.1.10").get("ports", []), "client ephemeral port listed as a device port")


def check_benign_scada_master(r: Result, out: RunOutput) -> None:
    expect_benign(r, out, "MODERATE", NOISE + ["OT protocol scanning", "Recon-then-write", "register scanning", "replay"])
    require(r, out, "IEC 104 control ASDU observed", src="10.20.0.5", dst="10.20.0.12")
    forbid(r, out, "IEC 104 control ASDU observed", src="10.20.0.12")
    require_heading(r, out, "IEC 60870-5-104")


def check_benign_office_lan(r: Result, out: RunOutput) -> None:
    expect_benign(r, out, "LOW", NOISE + ["Dual-zone", "ONVIF", "Beaconing", "SMB observed", "SNMPv1/v2 observed"])
    r.check("IT" in out.device("10.30.0.50").get("inferred_role", ""), "Windows PC not classified as an IT device")


def check_windows_hmi_dual_zone(r: Result, out: RunOutput) -> None:
    require(r, out, "Dual-zone", src="10.90.1.10")
    for finding in out.findings:
        r.check(finding["severity"] != "INTRUSION", f"unexpected INTRUSION: {finding['title']}")


def check_attack_creds_bruteforce(r: Result, out: RunOutput) -> None:
    for proto in ("FTP", "HTTP Basic", "HTTP Form", "Telnet", "MQTT", "SNMPv1/v2c"):
        require(r, out, f"Cleartext credentials extracted ({proto})")
    for proto in ("FTP", "HTTP", "Telnet"):
        require(r, out, f"Authentication failure burst ({proto})", src="10.40.0.99")
    require(r, out, "Obsolete TLS/SSL version observed (TLS 1.0)")
    require(r, out, "Weak TLS cipher observed (RC4")
    require(r, out, "SNMP community string 'plantRW' observed")
    require(r, out, "MQTT connection without credentials", src="10.40.0.31")
    forbid(r, out, "MQTT connection without credentials", src="10.40.0.30")
    for title in NOISE:
        forbid(r, out, title)


def check_attack_modbus_recon_injection(r: Result, out: RunOutput) -> None:
    require(r, out, "Possible Modbus register scanning", src="10.70.0.66", dst="10.70.0.20")
    forbid(r, out, "Possible Modbus register scanning", src="10.70.0.10")
    r.check(bool(out.involving("Duplicate or unsolicited Modbus responses", "10.70.0.20")),
            "replayed Modbus responses were not detected")
    require(r, out, "Possible OT protocol scanning", src="10.70.0.66")
    forbid(r, out, "Possible OT protocol scanning", src="10.70.0.10")
    require(r, out, "ICMP sweep", src="10.70.0.66")
    require(r, out, "Recon-then-write pattern observed", src="10.70.0.66")
    require(r, out, "Modbus write operation observed", src="10.70.0.66", dst="10.70.0.20")


def check_attack_arp_mitm(r: Result, out: RunOutput) -> None:
    require(r, out, "ARP-based Adversary-in-the-Middle attack pattern detected", dst="10.80.0.20", severity="INTRUSION")
    mitm = out.matching("ARP-based Adversary-in-the-Middle attack pattern detected")
    r.check(bool(mitm) and any("00:0c:29:aa:bb:cc" in line for line in mitm[0]["evidence"]),
            "spoofed MAC 00:0c:29:aa:bb:cc missing from the MITM evidence")
    forbid(r, out, "IP address claimed by multiple MACs in ARP")


def check_av_camera_voip(r: Result, out: RunOutput) -> None:
    require(r, out, "Active video stream requested via RTSP")
    require(r, out, "VoIP call signaling (SIP) observed")
    require(r, out, "Audio/voice streaming (RTP) observed")
    r.check(bool(out.involving("RTSP stream without authentication challenge", "10.50.0.11")),
            "unauthenticated camera 10.50.0.11 not flagged")
    r.check(not out.involving("RTSP stream without authentication challenge", "10.50.0.10"),
            "camera 10.50.0.10 issued a 401 challenge but was flagged as unauthenticated")
    require_heading(r, out, "Audio/Video Streaming")


def check_s7_engineering(r: Result, out: RunOutput) -> None:
    forbid(r, out, "IEC 61850 MMS without TLS")
    r.check(out.report.get("site") == "R&D Plant <North>", "site name with special characters was not preserved")
    r.check(len(out.svgs) >= 1, "no Purdue SVG written")


def check_arp_same_ms(r: Result, out: RunOutput) -> None:
    r.check(" 0 duplicate packets skipped" in out.log, "distinct ARP frames were counted as duplicates")


def check_modbus_value_jump(r: Result, out: RunOutput) -> None:
    require(r, out, "Modbus register value anomaly detected", src="10.11.0.20", dst="10.11.0.10")
    forbid(r, out, "Modbus write operation observed")


def check_attack_mitm_intercept(r: Result, out: RunOutput) -> None:
    mitm = out.matching("ARP-based Adversary-in-the-Middle attack pattern detected — traffic intercepted",
                        src="10.61.0.80", severity="INTRUSION")
    r.check(bool(mitm), "intercepted MITM not reported as INTRUSION from the impostor's own address 10.61.0.80")
    if mitm:
        finding = mitm[0]
        r.check({"10.61.0.70", "10.61.0.250"} <= set(finding["destination_ips"]),
                "MITM finding does not name the impersonated HMI and the poisoned PLC")
        r.check(finding["category"] == "Intrusion Indicators", f"MITM filed under {finding['category']}")
        r.check(any("00:0c:29:e6:14:0d" in line and "Impostor MAC" in line for line in finding["evidence"]),
                "impostor MAC missing from the evidence")
        try:
            span = float(finding["last_seen"]) - float(finding["first_seen"])
        except (TypeError, ValueError):
            span = -1
        r.check(15 <= span <= 30, f"attack window is {span:.1f} s, expected about 20 s")
    for title in ("Duplicate or unsolicited Modbus responses", "TCP reset storm observed",
                  "IP address claimed by multiple MACs", "Multiple MAC addresses observed for one IP",
                  "ARP sender MAC differs", "ARP replies without matching requests"):
        forbid(r, out, title)
    require(r, out, "Modbus write operation observed", src="10.61.0.251", dst="10.61.0.250")
    devices = out.report.get("devices", {})
    r.check("10.61.0.80" in devices, "impostor's own address 10.61.0.80 missing from the device inventory")
    for group in ("10.61.0.255", "224.0.0.251"):
        r.check(group not in devices, f"group address {group} listed as a device")
    r.check(out.band() in ("HIGH", "SEVERE"), f"risk band {out.band()} for a confirmed intrusion")
    require_heading(r, out, "Intrusion Indicators")
    # Readable layout: remediation plan, summary box, numbered steps.
    require_heading(r, out, "Remediation Plan")
    for text in ("What happened", "What to do now", "1. Contain now:", "3. Verify:", "OT safety note:",
                 "Isolate the impostor device and turn on ARP protection", "whose own address is 10.61.0.80,"):
        require_text(r, out, text)
    plan = out.report.get("remediation_plan", [])
    r.check(bool(plan) and plan[0]["priority"] == 1 and "impostor" in plan[0]["action"],
            "remediation plan does not open with isolating the impostor")
    r.check(all(f.get("report_id") for f in out.findings
                if f["category"] not in ("Protection Guidance", "Assessor Notes")),
            "JSON findings are missing their report IDs")


def check_truncated(r: Result, out: RunOutput) -> None:
    pcaps = out.report.get("pcaps", {})
    r.check(bool(pcaps), "truncated capture missing from the report's processed-capture list")
    r.check(any(p.get("truncated") for p in pcaps.values()), "truncated capture not flagged as truncated")
    r.check("Failed to inventory" not in out.log, "truncated capture reported as an inventory failure")


def check_dnp3_sample(r: Result, out: RunOutput) -> None:
    for title in ("DNP3 Cold Restart command observed", "DNP3 Warm Restart command observed",
                  "DNP3 Disable Spontaneous Messages observed", "DNP3 traffic without Secure Authentication evidence"):
        require(r, out, title)
    for title in NOISE:
        forbid(r, out, title)
    r.check(len(out.allowlist) == 3, f"flow allowlist has {len(out.allowlist)} rows, expected 3 real flows")
    # Two masters sent the same command: each pair is listed on its own.
    require_text(r, out, "Occurrences (2):")
    for pair in ("10.0.0.8 → 10.0.0.3", "10.0.0.9 → 10.0.0.3"):
        require_text(r, out, pair)
    restart_rows = [row for row in out.report.get("remediation_plan", [])
                    if row["action"].startswith("Confirm who stopped")]
    r.check(len(restart_rows) == 1 and len(restart_rows[0]["findings"]) == 3 and restart_rows[0]["priority"] == 1,
            "restart and silencing findings are not combined into one priority-1 plan row")


def check_multi_capture_out_of_order(r: Result, out: RunOutput) -> None:
    r.check(len(out.report.get("pcaps", {})) == 2, "both capture files should be analyzed")
    forbid(r, out, "Repetitive single-register polling pattern", src="10.12.0.10")
    hammer = out.matching("Repetitive single-register polling pattern", src="10.12.0.66")
    r.check(bool(hammer) and "about 50 requests per second" in hammer[0]["description"],
            "hammering client not reported at about 50 requests per second")
    forbid(r, out, "No NTP traffic observed in capture")
    gaps = out.matching("Capture discontinuity")
    r.check(any("plant_a.pcap" in g["description"] and "3 days" in g["description"] for g in gaps),
            "three-day pause in plant_a.pcap not reported in days")


CHECKS: Dict[str, Callable[[Result, RunOutput], None]] = {
    "benign_modbus_hmi": check_benign_modbus_hmi,
    "benign_scada_master": check_benign_scada_master,
    "benign_office_lan": check_benign_office_lan,
    "windows_hmi_dual_zone": check_windows_hmi_dual_zone,
    "attack_creds_bruteforce": check_attack_creds_bruteforce,
    "attack_modbus_recon_injection": check_attack_modbus_recon_injection,
    "attack_arp_mitm": check_attack_arp_mitm,
    "av_camera_voip": check_av_camera_voip,
    "s7_engineering": check_s7_engineering,
    "arp_same_ms": check_arp_same_ms,
    "modbus_value_jump": check_modbus_value_jump,
    "attack_mitm_intercept": check_attack_mitm_intercept,
    "truncated_capture": check_truncated,
    "dnp3_sample": check_dnp3_sample,
    "multi_capture_out_of_order": check_multi_capture_out_of_order,
}
SITES = {"benign_modbus_hmi": None, "s7_engineering": "R&D Plant <North>"}
LOOSE = {"benign_modbus_hmi"}


def main() -> int:
    parser = argparse.ArgumentParser(description="OTscope detection regression suite")
    parser.add_argument("--only", nargs="*", help="run only these scenarios")
    parser.add_argument("--keep", action="store_true", help="keep the temporary workspaces")
    parser.add_argument("--otscope", type=Path, default=DEFAULT_OTSCOPE, help="path to otscope.py under test")
    args = parser.parse_args()
    if shutil.which("tshark") is None:
        print("tshark not found on PATH — cannot run the regression suite.")
        return 2
    names = list(CHECKS)
    if args.only:
        names = [n for n in names if n in args.only]
    root = Path(tempfile.mkdtemp(prefix="otscope-regress-"))
    home = root / "home"
    home.mkdir()
    results: List[Result] = []
    try:
        for name in names:
            result = Result(name)
            capture_dir = root / "captures" / name
            capture_dir.mkdir(parents=True, exist_ok=True)
            if name == "dnp3_sample":
                if not DNP3_SAMPLE.exists():
                    result.summary = "skipped (sample pcap not present)"
                    results.append(result)
                    continue
                captures = [DNP3_SAMPLE]
            elif name == "truncated_capture":
                captures = [capture_dir / "truncated_capture.pcap"]
                SCENARIOS["benign_scada_master"]().write(str(captures[0]), truncate_tail=20)
            else:
                built = SCENARIOS[name]()
                nets = built if isinstance(built, dict) else {f"{name}.pcap": built}
                captures = []
                for filename, net in nets.items():
                    captures.append(capture_dir / filename)
                    net.write(str(captures[-1]))
            site = SITES.get(name, name)
            out = run_otscope(args.otscope, root / name, captures, name in LOOSE, site, home)
            expect_common(result, out)
            if out.report:
                CHECKS[name](result, out)
            risk = out.report.get("risk", {})
            result.summary = f"{len(out.findings)} findings, score {risk.get('score', '?')} {risk.get('band', '?')}"
            results.append(result)
            status = "PASS" if not result.failures else "FAIL"
            print(f"[{status}] {name:<32} {result.summary}", flush=True)
            for failure in result.failures:
                print(f"         - {failure}", flush=True)
    finally:
        if args.keep:
            print(f"\nWorkspaces kept in {root}")
        else:
            shutil.rmtree(root, ignore_errors=True)
    failed = [r for r in results if r.failures]
    print(f"\n{len(results) - len(failed)}/{len(results)} scenarios passed.")
    return 1 if failed else 0


if __name__ == "__main__":
    raise SystemExit(main())
